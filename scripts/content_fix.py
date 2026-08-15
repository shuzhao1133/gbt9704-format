"""
content_fix.py —— gbt9704-format 扩展：明显错字/重复的"自动改+批注版逐条留痕"

设计原则（与本 skill 其余规则一致）：
- 只处理"高置信度、无歧义"的错字/重复问题，不做模糊判断；专有名词、政策术语、地名、
  简称一律不动，拿不准的一律跳过（宁可不改）。
- 每一处改动都必须同时体现在两个产出里：
  clean 版——直接改好；批注版——原文保留、加批注说明改了什么、为什么改。
  不允许"改了但批注版看不出来"的静默改动。
- 具体识别"哪里是错字/重复"仍需人工（或调用 Claude）通读判断，本脚本只负责按判断结果
  批量执行"改 clean 版 + 批注留痕"这两个动作，不包含判断逻辑本身。

用法：
  python3 content_fix.py 文稿.docx fixes.json \
      -clean-out 终稿-内容已修复.docx -review-out 批注版-内容留痕.docx

fixes.json 格式：
  [{"anchor": "运心中形成", "replacement": "运行中形成",
    "category": "错字", "note": "\"运心中\"疑为\"运行中\"之误"},
   {"anchor": "深化极简审批和项目全程服务，深化极简审批和项目全程服务",
    "replacement": "深化极简审批和项目全程服务",
    "category": "逻辑重复", "note": "同一表述在同一句内重复两次，已删除重复部分"}]
"""
import argparse
import json
from copy import deepcopy

from docx import Document
from docx.oxml.text.paragraph import CT_P
from docx.oxml.table import CT_Tbl
from docx.text.paragraph import Paragraph
from docx.table import Table
from docx.text.run import Run

AUTHOR_BY_CATEGORY = {
    '错字': '错字',
    '逻辑重复': '逻辑重复',
}
DEFAULT_AUTHOR = '内容已修复'


def iter_blocks(doc):
    parent = doc.element.body
    for child in parent.iterchildren():
        if isinstance(child, CT_P):
            yield Paragraph(child, doc)
        elif isinstance(child, CT_Tbl):
            tbl = Table(child, doc)
            for row in tbl.rows:
                for cell in row.cells:
                    for p in cell.paragraphs:
                        yield p


def all_paragraphs(doc):
    return list(iter_blocks(doc))


def find_nth(text, sub, n):
    i = -1
    for _ in range(n):
        i = text.find(sub, i + 1)
        if i < 0:
            return -1
    return i


def split_run(para, run, offset):
    el = deepcopy(run._element)
    run._element.addnext(el)
    new = Run(el, para)
    t = run.text
    run.text, new.text = t[:offset], t[offset:]
    return new


def runs_for_range(para, start, end):
    pos = 0
    targets = []
    for run in list(para.runs):
        l = len(run.text)
        rs, re_ = pos, pos + l
        if re_ <= start or rs >= end:
            pos = re_
            continue
        r = run
        if rs < start:
            r = split_run(para, r, start - rs)
            rs = start
        if re_ > end:
            split_run(para, r, end - rs)
        targets.append(r)
        pos = re_
    return targets


def locate(paragraphs, anchor, occurrence=1):
    for para in paragraphs:
        if anchor in para.text:
            i = find_nth(para.text, anchor, occurrence)
            if i < 0:
                i = para.text.find(anchor)
            return para, i
    return None, -1


def apply_clean_fixes(path, fixes, out):
    doc = Document(path)
    paragraphs = all_paragraphs(doc)
    ok, missed = 0, []
    for fx in fixes:
        para, i = locate(paragraphs, fx['anchor'], fx.get('occurrence', 1))
        if para is None:
            missed.append(fx)
            continue
        runs = runs_for_range(para, i, i + len(fx['anchor']))
        if not runs:
            missed.append(fx)
            continue
        runs[0].text = fx['replacement']
        for r in runs[1:]:
            r.text = ''
        ok += 1
    doc.save(out)
    print(f'已输出内容已修复的 clean 版：{out}（成功修改 {ok} 处，原文一字未改的是批注版）')
    if missed:
        print(f'!! {len(missed)} 处锚点未找到，未修改：')
        for fx in missed:
            print(f'   - anchor={fx["anchor"]!r}')
    return ok, missed


def apply_review_annotations(path, fixes, out):
    doc = Document(path)
    paragraphs = all_paragraphs(doc)
    ok, missed = 0, []
    for fx in fixes:
        para, i = locate(paragraphs, fx['anchor'], fx.get('occurrence', 1))
        if para is None:
            missed.append(fx)
            continue
        runs = runs_for_range(para, i, i + len(fx['anchor']))
        if not runs:
            missed.append(fx)
            continue
        category = fx.get('category', '内容')
        author = AUTHOR_BY_CATEGORY.get(category, DEFAULT_AUTHOR)
        comment_text = f"【内容已修复：{category}】{fx.get('note', '')}原文：{fx['anchor']!r} → 改为：{fx['replacement']!r}"
        doc.add_comment(runs=runs, text=comment_text, author=author, initials='审')
        ok += 1
    doc.save(out)
    print(f'已输出批注留痕版：{out}（成功批注 {ok} 处，原文一字未改）')
    if missed:
        print(f'!! {len(missed)} 处锚点未找到，未批注：')
        for fx in missed:
            print(f'   - anchor={fx["anchor"]!r}')
    return ok, missed


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('docx')
    ap.add_argument('fixes_json')
    ap.add_argument('-clean-out', required=True)
    ap.add_argument('-review-out', required=True)
    args = ap.parse_args()

    fixes = json.loads(open(args.fixes_json, encoding='utf-8').read())
    apply_clean_fixes(args.docx, fixes, args.clean_out)
    apply_review_annotations(args.docx, fixes, args.review_out)


if __name__ == '__main__':
    main()
